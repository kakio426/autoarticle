
from docx import Document
from docx.shared import RGBColor, Inches, Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import io
import json
import os
import datetime

from engines.constants import THEMES

class WordEngine:
    def __init__(self, theme_name, school_name):
        self.theme = THEMES.get(theme_name, THEMES["웜 & 플레이풀"])
        self.school_name = school_name

    def generate(self, articles):
        doc = Document()
        
        # 기본 스타일 설정
        style = doc.styles['Normal']
        font = style.font
        font.name = 'Malgun Gothic'
        font.size = Pt(11)
        
        # 표지 페이지
        self._add_cover_page(doc)
        
        # 각 기사를 새 페이지에 추가
        for art in articles:
            doc.add_page_break()
            self._add_article_page(doc, art)
            
        # 헤더/푸터 추가
        self._add_header_footer(doc)
            
        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        return buffer

    def _add_header_footer(self, doc):
        """모든 페이지에 헤더와 푸터 추가"""
        section = doc.sections[0]
        
        # 헤더
        header = section.header
        header_para = header.paragraphs[0]
        header_para.text = f"{self.school_name} 소식지"
        header_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in header_para.runs:
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(*self.theme["main"])
            run.font.bold = True
        
        # 구분선 추가
        header_para_line = header.add_paragraph()
        header_para_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = header_para_line.add_run("─" * 50)
        run.font.color.rgb = RGBColor(200, 200, 200)
        run.font.size = Pt(8)
        
        # 푸터 (페이지 번호)
        footer = section.footer
        footer_para = footer.paragraphs[0]
        footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        # 페이지 번호 필드 추가
        run = footer_para.add_run()
        fldChar1 = OxmlElement('w:fldChar')
        fldChar1.set(qn('w:fldCharType'), 'begin')
        
        instrText = OxmlElement('w:instrText')
        instrText.set(qn('xml:space'), 'preserve')
        instrText.text = "PAGE"
        
        fldChar2 = OxmlElement('w:fldChar')
        fldChar2.set(qn('w:fldCharType'), 'end')
        
        run._r.append(fldChar1)
        run._r.append(instrText)
        run._r.append(fldChar2)
        
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(128, 128, 128)

    def _add_cover_page(self, doc):
        """전문적인 표지 페이지"""
        # 상단 여백
        for _ in range(8):
            doc.add_paragraph()
        
        # 학교명 (대제목)
        title = doc.add_heading(level=0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = title.add_run(self.school_name)
        run.font.size = Pt(44)
        run.font.color.rgb = RGBColor(*self.theme["main"])
        run.font.bold = True
        
        # 부제목
        now = datetime.datetime.now()
        subtitle = doc.add_paragraph()
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = subtitle.add_run(f"{now.year}학년도 {now.month}월 뉴스레터")
        run.font.size = Pt(20)
        run.font.color.rgb = RGBColor(100, 100, 100)
        
        # 장식 구분선
        doc.add_paragraph()
        deco = doc.add_paragraph()
        deco.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = deco.add_run("◆ ◆ ◆")
        run.font.size = Pt(16)
        run.font.color.rgb = RGBColor(*self.theme["accent"])
        
        # 발행일
        doc.add_paragraph()
        doc.add_paragraph()
        date_para = doc.add_paragraph()
        date_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = date_para.add_run(f"발행일: {now.strftime('%Y년 %m월 %d일')}")
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(120, 120, 120)

    def _add_article_page(self, doc, article):
        """개선된 기사 페이지 레이아웃"""
        
        # 기사 제목 (큰 헤딩)
        title = doc.add_heading(level=1)
        title.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = title.add_run(str(article.get('title', '')))
        run.font.size = Pt(22)
        run.font.color.rgb = RGBColor(*self.theme["main"])
        run.font.bold = True
        
        # 메타데이터 (표 형식으로 깔끔하게)
        meta_table = doc.add_table(rows=1, cols=1)
        meta_table.style = 'Light Grid Accent 1'
        
        cell = meta_table.cell(0, 0)
        
        # 셀 배경색
        shading_elm = OxmlElement('w:shd')
        shading_elm.set(qn('w:fill'), 'F5F5F5')
        cell._element.get_or_add_tcPr().append(shading_elm)
        
        # 메타 정보
        info_parts = []
        if article.get('date'): 
            info_parts.append(f"📅 일시: {article['date']}")
        if article.get('location'): 
            info_parts.append(f"📍 장소: {article['location']}")
        if article.get('grade'): 
            info_parts.append(f"👥 대상: {article['grade']}")
        
        p = cell.paragraphs[0]
        run = p.add_run("  |  ".join(info_parts))
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(80, 80, 80)
        
        doc.add_paragraph()  # 간격
        
        # 이미지 처리
        imgs_raw = article.get('images', '[]')
        imgs = json.loads(imgs_raw) if isinstance(imgs_raw, str) else imgs_raw
        valid_imgs = [p for p in imgs if os.path.exists(p)]
        
        if valid_imgs:
            img_count = min(len(valid_imgs), 2)
            
            if img_count == 1:
                # 단일 이미지: 중앙 정렬, 테두리 추가
                img_table = doc.add_table(rows=1, cols=1)
                img_table.alignment = WD_ALIGN_PARAGRAPH.CENTER
                
                cell = img_table.cell(0, 0)
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run()
                run.add_picture(valid_imgs[0], width=Inches(5.0))
                
                # 테두리 스타일
                tcPr = cell._element.get_or_add_tcPr()
                tcBorders = OxmlElement('w:tcBorders')
                for border_name in ['top', 'left', 'bottom', 'right']:
                    border = OxmlElement(f'w:{border_name}')
                    border.set(qn('w:val'), 'single')
                    border.set(qn('w:sz'), '12')
                    border.set(qn('w:color'), 'CCCCCC')
                    tcBorders.append(border)
                tcPr.append(tcBorders)
                
            else:
                # 2개 이미지: 좌우 배치
                img_table = doc.add_table(rows=1, cols=2)
                img_table.autofit = False
                img_table.allow_autofit = False
                
                for i in range(2):
                    cell = img_table.cell(0, i)
                    p = cell.paragraphs[0]
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    run = p.add_run()
                    run.add_picture(valid_imgs[i], width=Inches(3.2))
                    
                    # 셀 패딩
                    tcPr = cell._element.get_or_add_tcPr()
                    tcMar = OxmlElement('w:tcMar')
                    for margin in ['top', 'left', 'bottom', 'right']:
                        node = OxmlElement(f'w:{margin}')
                        node.set(qn('w:w'), '100')
                        node.set(qn('w:type'), 'dxa')
                        tcMar.append(node)
                    tcPr.append(tcMar)
        
        doc.add_paragraph()  # 간격
        
        # 본문 (단락 스타일 적용)
        content_para = doc.add_paragraph()
        content_para.paragraph_format.line_spacing = 1.5
        content_para.paragraph_format.space_after = Pt(12)
        
        run = content_para.add_run(str(article.get('content', '')))
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(40, 40, 40)
        
        # 기사 끝 구분선
        doc.add_paragraph()
        sep = doc.add_paragraph()
        sep.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = sep.add_run("• • •")
        run.font.color.rgb = RGBColor(*self.theme["accent"])
        run.font.size = Pt(14)
